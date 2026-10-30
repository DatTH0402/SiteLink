"""
reconcile.py - read-only API of the dump <-> web cell reconciliation dashboard.

The data is produced once a day by backend/reconcile_cells.py. For every technology
the LATEST SUCCESSFUL run is served.

  GET  /summary                 KPI cards, vendor / province breakdown, history
  GET  /items                   paged detail rows (server-side filter / sort)
  GET  /items/ids               ids of every row matching the filters (select-all)
  GET  /items/distinct/{col}    distinct values + counts (Excel-style header filter)
  POST /export                  Excel of the filtered rows (or of the selected ids)
  GET  /export-all              one workbook: summary + 6 sheets
"""
from __future__ import annotations

import io
import json
import re
import threading
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import false, func, or_
from sqlalchemy.orm import Session

from app.db.session import engine, get_db
from app.models.cell_reconcile import CellReconcileItem as Item, CellReconcileRun as Run
from app.services.list_query import (
    _check, _fold_py, _fold_sql, _norm, apply_column_filters, apply_sort,
)
from app.utils.deps import get_current_user

TECHS = ("3g", "4g", "5g")
DIRECTIONS = ("missing_in_web", "missing_in_dump")

_tables_ready = False
_tables_lock = threading.Lock()


def _ensure_tables() -> None:
    """The batch script creates the tables; this makes the API work even before its first run."""
    global _tables_ready
    if _tables_ready:
        return
    with _tables_lock:
        if _tables_ready:
            return
        from app.db.base import Base
        Base.metadata.create_all(bind=engine, tables=[Run.__table__, Item.__table__])
        _tables_ready = True


router = APIRouter(dependencies=[Depends(_ensure_tables)])

_EXPORT_SLOTS = threading.BoundedSemaphore(2)
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_EXPOSE = "X-Row-Count, Content-Disposition"

RUN_INT_FIELDS = [
    "dump_total", "dump_active", "dump_inactive", "web_total", "matched", "missing_in_web",
    "missing_in_dump", "missing_in_dump_uncovered", "dump_inactive_not_in_web",
    "new_missing_in_web", "new_missing_in_dump", "resolved_missing_in_web", "resolved_missing_in_dump",
]
ITEM_FIELDS = [
    "id", "tech", "direction", "cell_name", "subcell_id", "web_id", "site_name", "vendor", "mien",
    "tinh", "phuong_xa", "cell_status", "bbu_name", "rf", "oss", "dump_date", "cell_id", "node_id",
    "freq", "pci", "tac", "bandwidth", "mimo", "web_updated_at", "first_seen", "dump_covered",
]
TECH_LABELS = {
    "3g": {"node": None,        "freq": "UARFCN",    "pci": "PSC", "tac": "LAC", "bw": False},
    "4g": {"node": "EnodeB ID", "freq": "EARFCN",    "pci": "PCI", "tac": "TAC", "bw": True},
    "5g": {"node": "gNodeB ID", "freq": "SSB-ARFCN", "pci": "PCI", "tac": "TAC", "bw": True},
}
DIR_TITLE = {
    "missing_in_web":  "Chưa có trên Web",
    "missing_in_dump": "Không có trong Dump",
}


# ── helpers ──────────────────────────────────────────────────────────────────
def _tech(t: str) -> str:
    t = (t or "").lower()
    if t not in TECHS:
        raise HTTPException(400, "tech must be 3g, 4g or 5g")
    return t


def _direction(d: str) -> str:
    if d not in DIRECTIONS:
        raise HTTPException(400, "direction must be missing_in_web or missing_in_dump")
    return d


def _iso(dt) -> Optional[str]:
    return dt.isoformat() if dt else None


def _loads(s: Optional[str]):
    if not s:
        return None
    try:
        return json.loads(s)
    except ValueError:
        return None


def _latest_ok_run(db: Session, tech: str) -> Optional[Run]:
    return (db.query(Run).filter(Run.tech == tech, Run.status == "ok")
              .order_by(Run.id.desc()).first())


def _run_dict(r: Run, heavy: bool = True) -> Dict[str, Any]:
    d: Dict[str, Any] = {
        "id": r.id, "tech": r.tech, "batch_id": r.batch_id, "run_date": r.run_date,
        "started_at": _iso(r.started_at), "finished_at": _iso(r.finished_at),
        "status": r.status, "error": r.error, "prev_run_id": r.prev_run_id,
    }
    for f in RUN_INT_FIELDS:
        d[f] = int(getattr(r, f) or 0)
    if heavy:
        d["vendor_status"] = _loads(r.vendor_status)
        d["status_stats"] = _loads(r.status_stats)
    return d


def _age(first_seen: Optional[str], run_date: Optional[str]) -> Optional[int]:
    try:
        return (date.fromisoformat(run_date) - date.fromisoformat(first_seen)).days
    except Exception:
        return None


def _row(o: Item, run_date: Optional[str]) -> Dict[str, Any]:
    d = {f: getattr(o, f) for f in ITEM_FIELDS}
    d["age_days"] = _age(o.first_seen, run_date)
    return d


def _scoped(db: Session, tech: str, direction: str, search=None, vendor=None, mien=None,
            tinh=None, covered_only: bool = False, new_only: bool = False):
    """Rows of the latest successful run of `tech` / `direction` + the top-bar filters."""
    tech, direction = _tech(tech), _direction(direction)
    run = _latest_ok_run(db, tech)
    q = db.query(Item)
    if run is None:
        return q.filter(false()), None
    q = q.filter(Item.run_id == run.id, Item.direction == direction)
    if search and search.strip():
        like = "%" + search.strip() + "%"
        q = q.filter(or_(Item.cell_name.ilike(like), Item.site_name.ilike(like)))
    if vendor:
        q = q.filter(Item.vendor.in_(vendor))
    if mien:
        q = q.filter(Item.mien.in_(mien))
    if tinh:
        q = q.filter(Item.tinh.in_(tinh))
    if covered_only and direction == "missing_in_dump":
        q = q.filter(or_(Item.dump_covered.is_(None), Item.dump_covered != "N"))
    if new_only:
        q = q.filter(Item.first_seen == run.run_date)
    return q, run


def _sorted(q, sort_by: Optional[str], sort_dir: str):
    if sort_by:
        return apply_sort(q, Item, sort_by, sort_dir)
    return q.order_by(Item.cell_name.asc(), Item.id.asc())


# ── summary ──────────────────────────────────────────────────────────────────
@router.get("/summary")
def summary(db: Session = Depends(get_db), _=Depends(get_current_user)):
    techs: Dict[str, Any] = {}
    history: Dict[str, Any] = {}
    for tech in TECHS:
        run = _latest_ok_run(db, tech)
        last = db.query(Run).filter(Run.tech == tech).order_by(Run.id.desc()).first()
        entry: Dict[str, Any] = {
            "run": _run_dict(run) if run else None,
            "last_run": _run_dict(last, heavy=False) if last else None,
            "vendors": [], "provinces": [],
        }
        if run:
            vend = (db.query(Item.direction, Item.vendor, Item.dump_covered, func.count())
                      .filter(Item.run_id == run.id)
                      .group_by(Item.direction, Item.vendor, Item.dump_covered).all())
            entry["vendors"] = [
                {"direction": d, "vendor": v or "(Không rõ)", "covered": c != "N", "count": int(n)}
                for d, v, c, n in vend
            ]
            prov = (db.query(Item.direction, Item.tinh, func.count())
                      .filter(Item.run_id == run.id,
                              or_(Item.dump_covered.is_(None), Item.dump_covered != "N"))
                      .group_by(Item.direction, Item.tinh).all())
            agg: Dict[tuple, int] = {}
            for d, t, n in prov:
                key = (d, t or "(Chưa xác định)")
                agg[key] = agg.get(key, 0) + int(n)
            entry["provinces"] = [{"direction": d, "tinh": t, "count": n} for (d, t), n in agg.items()]
        techs[tech] = entry
        hist = (db.query(Run).filter(Run.tech == tech, Run.status == "ok")
                  .order_by(Run.id.desc()).limit(30).all())
        history[tech] = [
            {"run_date": h.run_date, "missing_in_web": int(h.missing_in_web or 0),
             "missing_in_dump": int(h.missing_in_dump or 0), "dump_active": int(h.dump_active or 0),
             "web_total": int(h.web_total or 0)}
            for h in reversed(hist)
        ]
    return {"techs": techs, "history": history, "generated_at": datetime.utcnow().isoformat() + "Z"}


# ── detail rows ──────────────────────────────────────────────────────────────
@router.get("/items")
def list_items(
    response: Response,
    tech: str,
    direction: str,
    skip: int = 0,
    limit: int = Query(500, ge=1, le=20000),
    search: Optional[str] = Query(None),
    vendor: Optional[List[str]] = Query(None),
    mien: Optional[List[str]] = Query(None),
    tinh: Optional[List[str]] = Query(None),
    covered_only: bool = False,
    new_only: bool = False,
    sort_by: Optional[str] = Query(None),
    sort_dir: str = Query("asc"),
    filters: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    q, run = _scoped(db, tech, direction, search, vendor, mien, tinh, covered_only, new_only)
    q = apply_column_filters(q, Item, filters)
    response.headers["X-Total-Count"] = str(q.count())
    q = _sorted(q, sort_by, sort_dir)
    rd = run.run_date if run else None
    return [_row(r, rd) for r in q.offset(skip).limit(limit).all()]


@router.get("/items/ids")
def list_ids(
    tech: str,
    direction: str,
    search: Optional[str] = Query(None),
    vendor: Optional[List[str]] = Query(None),
    mien: Optional[List[str]] = Query(None),
    tinh: Optional[List[str]] = Query(None),
    covered_only: bool = False,
    new_only: bool = False,
    filters: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    q, _run = _scoped(db, tech, direction, search, vendor, mien, tinh, covered_only, new_only)
    q = apply_column_filters(q, Item, filters)
    ids = [r[0] for r in q.with_entities(Item.id).order_by(Item.id).all()]
    return {"ids": ids, "total": len(ids)}


@router.get("/items/distinct/{column}")
def distinct_values(
    column: str,
    tech: str,
    direction: str,
    search: Optional[str] = Query(None),
    vendor: Optional[List[str]] = Query(None),
    mien: Optional[List[str]] = Query(None),
    tinh: Optional[List[str]] = Query(None),
    covered_only: bool = False,
    new_only: bool = False,
    filters: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    _check(Item, column)
    expr = _norm(Item, column)
    base, _run = _scoped(db, tech, direction, search, vendor, mien, tinh, covered_only, new_only)
    base = apply_column_filters(base, Item, filters, exclude=column)   # like Excel
    if q and q.strip():
        base = base.filter(_fold_sql(getattr(Item, column)).contains(_fold_py(q.strip()), autoescape=True))
    grouped = base.with_entities(expr.label("v"), func.count().label("n")).group_by(expr)
    rows = grouped.order_by(expr.asc().nulls_last()).limit(limit).all()
    total = db.query(func.count()).select_from(grouped.order_by(None).subquery()).scalar() or 0
    return {"values": [{"value": r.v, "count": r.n} for r in rows], "total_distinct": total}


# ── Excel export ─────────────────────────────────────────────────────────────
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xl(v: Any) -> Any:
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v


def _export_columns(tech: str, direction: str):
    L = TECH_LABELS[tech]
    cols = [("Cell Name", 30, "cell_name")]
    if direction == "missing_in_dump":
        cols += [("Subcell-ID", 14, "subcell_id"), ("Miền", 8, "mien"), ("Tỉnh", 22, "tinh"),
                 ("Phường/Xã", 24, "phuong_xa"), ("Site Name", 28, "site_name"),
                 ("Vendor", 12, "vendor"), ("Cell status (web)", 20, "cell_status")]
    else:
        cols += [("Vendor", 12, "vendor"), ("Trạng thái (dump)", 20, "cell_status"),
                 ("Miền (suy từ mã cell)", 12, "mien"), ("Tỉnh (suy từ mã cell)", 22, "tinh"),
                 ("Phường/Xã (suy từ mã cell)", 24, "phuong_xa")]
    cols += [("BBUname", 20, "bbu_name"), ("RF", 18, "rf"), ("OSS", 10, "oss"),
             ("Ngày dump", 14, "dump_date"), ("Cell ID", 14, "cell_id")]
    if L["node"]:
        cols.append((L["node"], 14, "node_id"))
    cols += [(L["freq"], 12, "freq"), (L["pci"], 10, "pci"), (L["tac"], 10, "tac")]
    if L["bw"]:
        cols.append(("Bandwidth", 12, "bandwidth"))
    cols.append(("MIMO", 10, "mimo"))
    if direction == "missing_in_dump":
        cols += [("Có nguồn dump", 18, "dump_covered_label"), ("Cập nhật web lần cuối", 20, "web_updated_at")]
    cols += [("Phát hiện lần đầu", 16, "first_seen"), ("Số ngày tồn tại", 14, "age_days")]
    return cols


def _val(o: Item, key: str, run_date: Optional[str]):
    if key == "dump_covered_label":
        return "Không (vendor chưa có nguồn dump)" if o.dump_covered == "N" else "Có"
    if key == "age_days":
        return _age(o.first_seen, run_date)
    return getattr(o, key, None)


def _fill_sheet(ws, cols, rows, run_date: Optional[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True, size=10)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    for i, (title, width, _k) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill, c.font = fill, font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    for o in rows:
        ws.append([_xl(_val(o, k, run_date)) for _t, _w, k in cols])
    ws.auto_filter.ref = "A1:%s1" % get_column_letter(len(cols))


def _stream(wb, filename: str, count: int) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.read()]), media_type=_XLSX,
        headers={"Content-Disposition": 'attachment; filename="%s"' % filename,
                 "X-Row-Count": str(count), "Access-Control-Expose-Headers": _EXPOSE})


def _pf(p: Dict[str, Any], k: str) -> Optional[str]:
    v = p.get(k)
    if isinstance(v, (list, tuple)):
        v = v[0] if v else None
    return str(v) if v not in (None, "") else None


def _pm(p: Dict[str, Any], k: str) -> List[str]:
    v = p.get(k)
    if v is None:
        return []
    if not isinstance(v, (list, tuple)):
        v = [v]
    return [str(x) for x in v if x not in (None, "")]


def _pb(p: Dict[str, Any], k: str) -> bool:
    v = p.get(k)
    if isinstance(v, (list, tuple)):
        v = v[0] if v else None
    return str(v).lower() in ("true", "1", "yes")


class ExportRequest(BaseModel):
    tech: str
    direction: str
    ids: Optional[List[int]] = None
    params: Dict[str, Any] = Field(default_factory=dict)
    filters: Optional[str] = None
    sort_by: Optional[str] = None
    sort_dir: str = "asc"


@router.post("/export")
def export_items(body: ExportRequest, db: Session = Depends(get_db), _=Depends(get_current_user)):
    tech, direction = _tech(body.tech), _direction(body.direction)
    with _EXPORT_SLOTS:
        p = body.params or {}
        if body.ids:
            run = _latest_ok_run(db, tech)
            q = db.query(Item).filter(Item.id.in_(list(body.ids)))
            q = q.filter(Item.run_id == run.id, Item.direction == direction) if run else q.filter(false())
        else:
            q, run = _scoped(db, tech, direction, _pf(p, "search"), _pm(p, "vendor"), _pm(p, "mien"),
                             _pm(p, "tinh"), _pb(p, "covered_only"), _pb(p, "new_only"))
            q = apply_column_filters(q, Item, body.filters)
        rows = _sorted(q, body.sort_by, body.sort_dir).all()
        rd = run.run_date if run else None
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = ("%s - %s" % (tech.upper(), DIR_TITLE[direction]))[:31]
        _fill_sheet(ws, _export_columns(tech, direction), rows, rd)
        name = "Doi_soat_%s_%s_%s.xlsx" % (
            tech.upper(), "chua_co_tren_web" if direction == "missing_in_web" else "khong_co_trong_dump",
            rd or datetime.now().strftime("%Y-%m-%d"))
        return _stream(wb, name, len(rows))


@router.get("/export-all")
def export_all(db: Session = Depends(get_db), _=Depends(get_current_user)):
    with _EXPORT_SLOTS:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tổng hợp"
        head = ["Công nghệ", "Ngày đối soát", "Cell trong dump", "Cell active trong dump", "Cell trên web",
                "Active trong dump - chưa có trên web", "(mới)", "(đã xử lý)",
                "Có trên web - không có trong dump", "(mới)", "(đã xử lý)",
                "Web: vendor chưa có nguồn dump"]
        ws.append(head)
        fill = PatternFill("solid", fgColor="1F4E79")
        for i in range(1, len(head) + 1):
            c = ws.cell(row=1, column=i)
            c.fill, c.font = fill, Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = 24
        ws.column_dimensions["A"].width = 12
        ws.row_dimensions[1].height = 42
        total = 0
        runs = {}
        for tech in TECHS:
            r = _latest_ok_run(db, tech)
            runs[tech] = r
            if r:
                ws.append([tech.upper(), r.run_date, r.dump_total, r.dump_active, r.web_total,
                           r.missing_in_web, r.new_missing_in_web, r.resolved_missing_in_web,
                           r.missing_in_dump, r.new_missing_in_dump, r.resolved_missing_in_dump,
                           r.missing_in_dump_uncovered])
        for tech in TECHS:
            run = runs[tech]
            if not run:
                continue
            for direction in DIRECTIONS:
                rows = (db.query(Item).filter(Item.run_id == run.id, Item.direction == direction)
                          .order_by(Item.tinh.asc().nulls_last(), Item.cell_name.asc()).all())
                total += len(rows)
                sheet = wb.create_sheet(("%s - %s" % (tech.upper(), DIR_TITLE[direction]))[:31])
                _fill_sheet(sheet, _export_columns(tech, direction), rows, run.run_date)
        return _stream(wb, "Doi_soat_cell_tong_hop_%s.xlsx" % datetime.now().strftime("%Y-%m-%d"), total)
