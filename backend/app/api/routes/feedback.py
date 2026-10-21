"""
feedback.py – Support / feedback ticket API   (prefix: /api/v1/feedback)

Users  : create tickets, follow / reply to / close / reopen / rate their own tickets.
Admins : see everything, change status / priority / category / assignee / due date,
         write resolution, internal notes, bulk update, export, delete.
"""
import io
import mimetypes
import os
import re
import uuid
from datetime import date, datetime, time, timedelta, timezone
from typing import List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.models.feedback import FeedbackAttachment, FeedbackEvent, FeedbackTicket
from app.models.user import User, UserRole
from app.utils.audit import log_action
from app.utils.deps import get_current_user, require_admin

router = APIRouter()

VN_TZ = timezone(timedelta(hours=7))

STATUS_LABEL = {
    "new": "Mới", "in_progress": "Đang xử lý", "waiting_user": "Chờ phản hồi",
    "planned": "Đã lên kế hoạch", "resolved": "Đã xử lý", "closed": "Đã đóng", "rejected": "Từ chối",
}
OPEN_STATUSES = ("new", "in_progress", "waiting_user", "planned")
DONE_STATUSES = ("resolved", "closed", "rejected")
CATEGORY_LABEL = {
    "bug": "Báo lỗi", "data": "Sai/thiếu dữ liệu", "change": "Yêu cầu thay đổi",
    "feature": "Đề xuất tính năng", "help": "Cần hỗ trợ / hướng dẫn",
    "account": "Tài khoản / phân quyền", "other": "Khác",
}
PRIORITY_LABEL = {"low": "Thấp", "normal": "Bình thường", "high": "Cao", "urgent": "Khẩn cấp"}
PRIORITY_RANK = {"low": 1, "normal": 2, "high": 3, "urgent": 4}

UPLOAD_ROOT = os.environ.get("FEEDBACK_UPLOAD_DIR") or os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "..", "uploads", "feedback"))
os.makedirs(UPLOAD_ROOT, exist_ok=True)

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILES_PER_REQUEST = 10
MAX_FILES_PER_TICKET = 30
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
ALLOWED_EXT = IMAGE_EXT | {".pdf", ".txt", ".log", ".csv", ".xlsx", ".xls",
                           ".docx", ".doc", ".pptx", ".zip"}
DAILY_TICKET_LIMIT = 20          # anti-spam for normal users


# ── helpers ───────────────────────────────────────────────────────────────────
def _now():
    return datetime.now(timezone.utc)


def _iso(d):
    return d.isoformat() if d else None


def _is_admin(u: User) -> bool:
    return u.role == UserRole.admin


def _uname(u: User) -> str:
    return u.full_name or u.username


def _code(t: FeedbackTicket) -> str:
    return f"FB-{t.id:05d}"


def _clip(s, n=480):
    if s is None:
        return None
    s = str(s)
    return s if len(s) <= n else s[: n - 1] + "…"


def _check(value, allowed, field):
    if value not in allowed:
        raise HTTPException(status_code=400, detail=f"Giá trị '{field}' không hợp lệ: {value}")
    return value


def _safe_name(name: Optional[str]) -> str:
    n = os.path.basename((name or "file").replace("\\", "/"))
    root, ext = os.path.splitext(n)
    root = re.sub(r"[^\w.\- ()\[\]]", "_", root).strip(" .") or "file"
    ext = re.sub(r"[^\w.]", "", ext)
    return root[:120] + ext[:12]


def _get_ticket(db: Session, ticket_id: int, user: User) -> FeedbackTicket:
    t = db.get(FeedbackTicket, ticket_id)
    if not t or (not _is_admin(user) and t.created_by != user.id):
        raise HTTPException(status_code=404, detail="Không tìm thấy yêu cầu")
    return t


def _event(db, t, user, kind, body=None, old=None, new=None, internal=False):
    ev = FeedbackEvent(
        ticket_id=t.id,
        user_id=user.id if user else None,
        user_name=_uname(user) if user else "Hệ thống",
        user_role=user.role.value if user else "system",
        kind=kind, body=body, old_value=_clip(old), new_value=_clip(new),
        is_internal=internal,
    )
    db.add(ev)
    t.last_activity_at = _now()
    return ev


def _set_status(db, t, new_status, user, body=None):
    old = t.status
    if old == new_status:
        return False
    now = _now()
    t.status = new_status
    if new_status == "resolved":
        t.resolved_at, t.closed_at = now, None
    elif new_status == "rejected":
        t.resolved_at, t.closed_at = now, None
    elif new_status == "closed":
        t.closed_at = now
        t.resolved_at = t.resolved_at or now
    else:
        t.resolved_at, t.closed_at = None, None
    _event(db, t, user, "status", body=body, old=old, new=new_status)
    return True


# ── serialisers ───────────────────────────────────────────────────────────────
def _ticket_dict(t: FeedbackTicket, user: User) -> dict:
    admin = _is_admin(user)
    overdue = bool(t.due_date and t.status in OPEN_STATUSES and t.due_date < _now())
    return {
        "id": t.id, "code": _code(t), "title": t.title,
        "category": t.category, "priority": t.priority, "status": t.status,
        "related_ref": t.related_ref,
        "created_by": t.created_by, "created_by_name": t.created_by_name,
        "created_by_email": t.created_by_email if admin else None,
        "assignee_id": t.assignee_id, "assignee_name": t.assignee_name,
        "due_date": _iso(t.due_date), "overdue": overdue,
        "resolution": t.resolution,
        "rating": t.rating, "rating_comment": t.rating_comment,
        "reopen_count": t.reopen_count or 0, "comment_count": 0,
        "created_at": _iso(t.created_at), "updated_at": _iso(t.updated_at),
        "last_activity_at": _iso(t.last_activity_at),
        "first_response_at": _iso(t.first_response_at),
        "resolved_at": _iso(t.resolved_at), "closed_at": _iso(t.closed_at),
        "is_owner": t.created_by == user.id,
    }


def _event_dict(e: FeedbackEvent) -> dict:
    return {
        "id": e.id, "kind": e.kind, "body": e.body,
        "old_value": e.old_value, "new_value": e.new_value,
        "is_internal": bool(e.is_internal),
        "user_id": e.user_id, "user_name": e.user_name, "user_role": e.user_role,
        "created_at": _iso(e.created_at),
    }


def _att_dict(a: FeedbackAttachment) -> dict:
    ext = os.path.splitext(a.filename or "")[1].lower()
    return {
        "id": a.id, "filename": a.filename, "content_type": a.content_type,
        "size": a.size or 0, "event_id": a.event_id,
        "uploaded_by": a.uploaded_by, "uploaded_by_name": a.uploaded_by_name,
        "created_at": _iso(a.created_at), "is_image": ext in IMAGE_EXT,
    }


def _detail(t: FeedbackTicket, user: User) -> dict:
    admin = _is_admin(user)
    hidden = {e.id for e in t.events if e.is_internal}
    d = _ticket_dict(t, user)
    d.update({
        "description": t.description,
        "page_url": t.page_url,
        "user_agent": t.user_agent if admin else None,
        "events": [_event_dict(e) for e in t.events if admin or not e.is_internal],
        "attachments": [_att_dict(a) for a in t.attachments
                        if admin or a.event_id not in hidden],
    })
    d["comment_count"] = sum(1 for e in t.events if e.kind == "comment")
    return d


# ── filters / sorting ─────────────────────────────────────────────────────────
class TicketFilters:
    def __init__(
        self,
        status: Optional[List[str]] = Query(None),
        category: Optional[List[str]] = Query(None),
        priority: Optional[List[str]] = Query(None),
        assignee_id: Optional[int] = Query(None),
        unassigned: Optional[bool] = Query(None),
        assigned_to_me: Optional[bool] = Query(None),
        mine: Optional[bool] = Query(None),
        overdue: Optional[bool] = Query(None),
        scope: Optional[str] = Query(None),            # open | done
        search: Optional[str] = Query(None),
        date_from: Optional[date] = Query(None),
        date_to: Optional[date] = Query(None),
    ):
        self.status, self.category, self.priority = status, category, priority
        self.assignee_id, self.unassigned = assignee_id, unassigned
        self.assigned_to_me, self.mine, self.overdue = assigned_to_me, mine, overdue
        self.scope, self.search = scope, search
        self.date_from, self.date_to = date_from, date_to


_T = FeedbackTicket
_SORT_COLS = {
    "id": _T.id, "created_at": _T.created_at, "updated_at": _T.updated_at,
    "last_activity_at": _T.last_activity_at, "title": _T.title, "status": _T.status,
    "category": _T.category, "due_date": _T.due_date,
    "created_by_name": _T.created_by_name, "assignee_name": _T.assignee_name,
}


def _filtered(db: Session, user: User, f: TicketFilters):
    admin = _is_admin(user)
    q = db.query(_T)
    if not admin or f.mine:
        q = q.filter(_T.created_by == user.id)
    if f.status:
        q = q.filter(_T.status.in_(f.status))
    if f.category:
        q = q.filter(_T.category.in_(f.category))
    if f.priority:
        q = q.filter(_T.priority.in_(f.priority))
    if admin:
        if f.unassigned:
            q = q.filter(_T.assignee_id.is_(None))
        if f.assigned_to_me:
            q = q.filter(_T.assignee_id == user.id)
        if f.assignee_id:
            q = q.filter(_T.assignee_id == f.assignee_id)
    if f.scope == "open":
        q = q.filter(_T.status.in_(OPEN_STATUSES))
    elif f.scope == "done":
        q = q.filter(_T.status.in_(DONE_STATUSES))
    if f.overdue:
        q = q.filter(_T.due_date.isnot(None), _T.due_date < _now(), _T.status.in_(OPEN_STATUSES))
    if f.date_from:
        q = q.filter(_T.created_at >= datetime.combine(f.date_from, time.min, tzinfo=VN_TZ))
    if f.date_to:
        q = q.filter(_T.created_at < datetime.combine(f.date_to + timedelta(days=1), time.min, tzinfo=VN_TZ))
    if f.search and f.search.strip():
        s = f.search.strip()
        like = f"%{s}%"
        conds = [_T.title.ilike(like), _T.description.ilike(like), _T.related_ref.ilike(like),
                 _T.created_by_name.ilike(like), _T.assignee_name.ilike(like)]
        m = re.fullmatch(r"(?i)(?:fb-?)?0*(\d+)", s)
        if m:
            conds.append(_T.id == int(m.group(1)))
        q = q.filter(or_(*conds))
    return q


def _sorted(q, sort_by: Optional[str], sort_dir: str):
    if sort_by == "priority":
        col = case(*[(_T.priority == k, v) for k, v in PRIORITY_RANK.items()], else_=0)
    else:
        col = _SORT_COLS.get(sort_by or "", _T.last_activity_at)
    col = (col.asc() if sort_dir == "asc" else col.desc()).nulls_last()
    return q.order_by(col, _T.id.desc())


# ── apply changes (shared by PATCH and bulk-update) ──────────────────────────
def _apply_changes(db: Session, t: FeedbackTicket, user: User, data: dict, note: str = ""):
    admin = _is_admin(user)
    changed = []

    for f in ("title", "description", "related_ref"):
        if f in data:
            new = (data[f] or "").strip() or None
            if f in ("title", "description") and not new:
                raise HTTPException(400, f"'{f}' không được để trống")
            old = getattr(t, f)
            if new != old:
                setattr(t, f, new)
                _event(db, t, user, "edit", body=f, old=_clip(old, 200), new=_clip(new, 200))
                changed.append(f)

    if "category" in data:
        new = _check(data["category"], CATEGORY_LABEL, "category")
        if new != t.category:
            _event(db, t, user, "category", old=t.category, new=new)
            t.category = new
            changed.append("category")

    if admin and "priority" in data:
        new = _check(data["priority"], PRIORITY_LABEL, "priority")
        if new != t.priority:
            _event(db, t, user, "priority", old=t.priority, new=new)
            t.priority = new
            changed.append("priority")

    if admin and "assignee_id" in data:
        aid = data["assignee_id"]
        if aid is None:
            if t.assignee_id is not None:
                _event(db, t, user, "assign", old=t.assignee_name, new="")
                t.assignee_id, t.assignee_name = None, None
                changed.append("assignee_id")
        elif aid != t.assignee_id:
            a = db.get(User, aid)
            if not a or not a.is_active or a.role != UserRole.admin:
                raise HTTPException(400, "Người xử lý phải là quản trị viên đang hoạt động")
            _event(db, t, user, "assign", old=t.assignee_name, new=_uname(a))
            t.assignee_id, t.assignee_name = a.id, _uname(a)
            changed.append("assignee_id")
            if t.status == "new" and "status" not in data:
                _set_status(db, t, "in_progress", None, body="Tự động: đã có người nhận xử lý")

    if admin and "status" in data:
        new = _check(data["status"], STATUS_LABEL, "status")
        if _set_status(db, t, new, user, body=note or None):
            changed.append("status")

    if admin and "due_date" in data:
        new = data["due_date"]
        if _iso(new) != _iso(t.due_date):
            _event(db, t, user, "due", old=_iso(t.due_date), new=_iso(new))
            t.due_date = new
            changed.append("due_date")

    if admin and "resolution" in data:
        new = (data["resolution"] or "").strip() or None
        if new != t.resolution:
            t.resolution = new
            if new:
                _event(db, t, user, "resolution", body=new)
            changed.append("resolution")

    if admin and changed and t.first_response_at is None and t.created_by != user.id:
        t.first_response_at = _now()
    return changed


# ── payloads ──────────────────────────────────────────────────────────────────
class TicketCreate(BaseModel):
    title: str = Field(..., min_length=3, max_length=255)
    description: str = Field(..., min_length=5, max_length=20000)
    category: str = "help"
    priority: str = "normal"
    related_ref: Optional[str] = Field(None, max_length=255)
    page_url: Optional[str] = Field(None, max_length=500)
    user_agent: Optional[str] = Field(None, max_length=500)


class TicketUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    category: Optional[str] = None
    priority: Optional[str] = None
    related_ref: Optional[str] = None
    status: Optional[str] = None
    assignee_id: Optional[int] = None
    due_date: Optional[datetime] = None
    resolution: Optional[str] = None
    note: Optional[str] = None


class CommentCreate(BaseModel):
    body: str = Field(..., min_length=1, max_length=20000)
    is_internal: bool = False


class ReasonPayload(BaseModel):
    reason: Optional[str] = Field(None, max_length=2000)


class RatePayload(BaseModel):
    rating: int = Field(..., ge=1, le=5)
    comment: Optional[str] = Field(None, max_length=2000)


class BulkUpdatePayload(BaseModel):
    ids: List[int]
    changes: dict


class BulkDeletePayload(BaseModel):
    ids: List[int]


# ═════════════════════════════════════════════════════════════════════════════
# STATIC ROUTES (must be declared before /{ticket_id})
# ═════════════════════════════════════════════════════════════════════════════
@router.get("/stats")
def stats(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    admin = _is_admin(user)
    base = db.query(_T)
    if not admin:
        base = base.filter(_T.created_by == user.id)
    now = _now()

    def grouped(col, q=None):
        qq = (q if q is not None else base).with_entities(col, func.count(_T.id)).group_by(col)
        return {k: v for k, v in qq.all()}

    by_status = grouped(_T.status)
    rows = base.filter(_T.created_at >= now - timedelta(days=90)).with_entities(
        _T.created_at, _T.first_response_at, _T.resolved_at).all()
    fr = [(b - a).total_seconds() for a, b, _ in rows if b]
    rs = [(c - a).total_seconds() for a, _, c in rows if c]

    def avg_h(xs):
        return round(sum(xs) / len(xs) / 3600, 1) if xs else None

    return {
        "total": sum(by_status.values()),
        "open": sum(v for k, v in by_status.items() if k in OPEN_STATUSES),
        "by_status": by_status,
        "by_category": grouped(_T.category),
        "by_priority_open": grouped(_T.priority, base.filter(_T.status.in_(OPEN_STATUSES))),
        "unassigned": base.filter(_T.assignee_id.is_(None), _T.status.in_(OPEN_STATUSES)).count(),
        "overdue": base.filter(_T.due_date.isnot(None), _T.due_date < now,
                               _T.status.in_(OPEN_STATUSES)).count(),
        "avg_first_response_h": avg_h(fr),
        "avg_resolution_h": avg_h(rs),
    }


@router.get("/badge")
def badge(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """admin: number of NEW tickets;  user: own tickets waiting for the user's reply."""
    if _is_admin(user):
        n = db.query(_T).filter(_T.status == "new").count()
    else:
        n = db.query(_T).filter(_T.created_by == user.id, _T.status == "waiting_user").count()
    return {"count": n}


@router.get("/assignees")
def assignees(db: Session = Depends(get_db), _=Depends(require_admin)):
    rows = (db.query(User).filter(User.role == UserRole.admin, User.is_active.is_(True))
            .order_by(User.full_name, User.username).all())
    return [{"id": u.id, "name": _uname(u)} for u in rows]


def _xl(v):
    """neutralise spreadsheet formula injection"""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _vn(d):
    return d.astimezone(VN_TZ).strftime("%Y-%m-%d %H:%M") if d else ""


@router.get("/export")
def export_tickets(
    f: TicketFilters = Depends(),
    sort_by: Optional[str] = Query(None),
    sort_dir: str = Query("desc"),
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = _sorted(_filtered(db, user, f), sort_by, sort_dir).limit(20000).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"
    head = ["Mã", "Tiêu đề", "Loại", "Ưu tiên", "Trạng thái", "Người gửi", "Email", "Người xử lý",
            "Tạo lúc", "Cập nhật", "Hạn xử lý", "Phản hồi đầu", "Đã xử lý lúc", "Kết quả xử lý",
            "Site/Cell liên quan", "Trang", "Mở lại", "Đánh giá", "Mô tả"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1677FF")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for t in rows:
        ws.append([_xl(x) for x in [
            _code(t), t.title, CATEGORY_LABEL.get(t.category, t.category),
            PRIORITY_LABEL.get(t.priority, t.priority), STATUS_LABEL.get(t.status, t.status),
            t.created_by_name or "", t.created_by_email or "", t.assignee_name or "",
            _vn(t.created_at), _vn(t.updated_at), _vn(t.due_date), _vn(t.first_response_at),
            _vn(t.resolved_at), t.resolution or "", t.related_ref or "", t.page_url or "",
            t.reopen_count or 0, t.rating or "", t.description or "",
        ]])
    widths = [11, 40, 22, 14, 16, 22, 28, 22, 17, 17, 17, 17, 17, 40, 28, 28, 8, 9, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"Feedback_{datetime.now(VN_TZ):%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/")
def list_tickets(
    response: Response,
    f: TicketFilters = Depends(),
    sort_by: Optional[str] = Query(None),
    sort_dir: str = Query("desc"),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=1000),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    q = _filtered(db, user, f)
    response.headers["X-Total-Count"] = str(q.count())
    rows = _sorted(q, sort_by, sort_dir).offset(skip).limit(limit).all()
    ids = [r.id for r in rows]
    counts = {}
    if ids:
        counts = dict(
            db.query(FeedbackEvent.ticket_id, func.count(FeedbackEvent.id))
            .filter(FeedbackEvent.ticket_id.in_(ids), FeedbackEvent.kind == "comment")
            .group_by(FeedbackEvent.ticket_id).all())
    out = []
    for r in rows:
        d = _ticket_dict(r, user)
        d["comment_count"] = counts.get(r.id, 0)
        out.append(d)
    return out


@router.post("/", status_code=201)
def create_ticket(
    payload: TicketCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    cat = _check(payload.category, CATEGORY_LABEL, "category")
    pri = _check(payload.priority, PRIORITY_LABEL, "priority")
    if not _is_admin(user):
        recent = db.query(_T).filter(_T.created_by == user.id,
                                     _T.created_at >= _now() - timedelta(days=1)).count()
        if recent >= DAILY_TICKET_LIMIT:
            raise HTTPException(429, "Bạn đã gửi quá nhiều yêu cầu trong 24 giờ qua")
    t = FeedbackTicket(
        title=payload.title.strip(), description=payload.description.strip(),
        category=cat, priority=pri, status="new",
        related_ref=(payload.related_ref or "").strip() or None,
        page_url=payload.page_url, user_agent=payload.user_agent,
        created_by=user.id, created_by_name=_uname(user), created_by_email=user.email,
        last_activity_at=_now(),
    )
    db.add(t)
    db.commit()
    db.refresh(t)
    try:
        log_action(db, user, "CREATE", "feedback_tickets", t.id,
                   new_value={"title": t.title, "category": cat, "priority": pri})
    except Exception:
        pass
    return _detail(t, user)


@router.post("/bulk-update")
def bulk_update(payload: BulkUpdatePayload, db: Session = Depends(get_db),
                user: User = Depends(require_admin)):
    allowed = {"status", "priority", "assignee_id", "category"}
    data = {k: v for k, v in payload.changes.items() if k in allowed}
    if not payload.ids or not data:
        raise HTTPException(400, "Thiếu danh sách yêu cầu hoặc nội dung thay đổi")
    updated, errors = 0, []
    for tid in payload.ids[:500]:
        t = db.get(_T, tid)
        if not t:
            errors.append(f"#{tid}: không tồn tại")
            continue
        try:
            if _apply_changes(db, t, user, data):
                updated += 1
            db.commit()
        except HTTPException as e:
            db.rollback()
            errors.append(f"{_code(t)}: {e.detail}")
        except Exception as e:  # noqa
            db.rollback()
            errors.append(f"#{tid}: {e}")
    return {"updated": updated, "errors": errors}


@router.post("/bulk-delete")
def bulk_delete(payload: BulkDeletePayload, db: Session = Depends(get_db),
                user: User = Depends(require_admin)):
    deleted = 0
    for tid in payload.ids[:500]:
        t = db.get(_T, tid)
        if not t:
            continue
        files = [a.stored_name for a in t.attachments]
        db.delete(t)
        db.commit()
        for fn in files:
            try:
                os.remove(os.path.join(UPLOAD_ROOT, fn))
            except OSError:
                pass
        try:
            log_action(db, user, "DELETE", "feedback_tickets", tid)
        except Exception:
            pass
        deleted += 1
    return {"deleted": deleted}


# ── attachments (static prefix, before /{ticket_id}) ─────────────────────────
@router.get("/attachments/{att_id}/download")
def download_attachment(att_id: int, db: Session = Depends(get_db),
                        user: User = Depends(get_current_user)):
    a = db.get(FeedbackAttachment, att_id)
    if not a:
        raise HTTPException(404, "Không tìm thấy tệp")
    _get_ticket(db, a.ticket_id, user)                       # 404 if not allowed
    if not _is_admin(user) and a.event_id:
        ev = db.get(FeedbackEvent, a.event_id)
        if ev and ev.is_internal:
            raise HTTPException(404, "Không tìm thấy tệp")
    path = os.path.join(UPLOAD_ROOT, a.stored_name)
    if not os.path.isfile(path):
        raise HTTPException(404, "Tệp không còn trên máy chủ")
    return FileResponse(
        path, filename=a.filename,
        media_type=a.content_type or "application/octet-stream",
        headers={"X-Content-Type-Options": "nosniff"})


@router.delete("/attachments/{att_id}")
def delete_attachment(att_id: int, db: Session = Depends(get_db),
                      user: User = Depends(get_current_user)):
    a = db.get(FeedbackAttachment, att_id)
    if not a:
        raise HTTPException(404, "Không tìm thấy tệp")
    _get_ticket(db, a.ticket_id, user)
    if not _is_admin(user) and a.uploaded_by != user.id:
        raise HTTPException(403, "Bạn chỉ có thể xóa tệp do mình tải lên")
    try:
        os.remove(os.path.join(UPLOAD_ROOT, a.stored_name))
    except OSError:
        pass
    db.delete(a)
    db.commit()
    return {"message": "Deleted"}


# ═════════════════════════════════════════════════════════════════════════════
# PER-TICKET ROUTES
# ═════════════════════════════════════════════════════════════════════════════
@router.get("/{ticket_id}")
def get_ticket(ticket_id: int, db: Session = Depends(get_db),
               user: User = Depends(get_current_user)):
    return _detail(_get_ticket(db, ticket_id, user), user)


@router.patch("/{ticket_id}")
def update_ticket(ticket_id: int, payload: TicketUpdate, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    t = _get_ticket(db, ticket_id, user)
    admin = _is_admin(user)
    data = payload.model_dump(exclude_unset=True)
    note = (data.pop("note", None) or "").strip()
    owner_fields = {"title", "description", "category", "related_ref"}
    allowed = (owner_fields | {"priority", "status", "assignee_id", "due_date", "resolution"}
               if admin else owner_fields)
    bad = [k for k in data if k not in allowed]
    if bad:
        raise HTTPException(403, f"Bạn không có quyền sửa: {', '.join(bad)}")
    if not admin and t.status != "new":
        raise HTTPException(403, "Chỉ có thể chỉnh sửa khi yêu cầu còn ở trạng thái 'Mới'")
    if "status" in data and data["status"] is None:
        raise HTTPException(400, "Trạng thái không hợp lệ")
    _apply_changes(db, t, user, data, note)
    db.commit()
    db.refresh(t)
    return _detail(t, user)


@router.post("/{ticket_id}/comments", status_code=201)
def add_comment(ticket_id: int, payload: CommentCreate, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    t = _get_ticket(db, ticket_id, user)
    admin = _is_admin(user)
    internal = bool(payload.is_internal) and admin
    if t.status == "closed":
        raise HTTPException(400, "Yêu cầu đã đóng – hãy mở lại để tiếp tục trao đổi")
    ev = _event(db, t, user, "note" if internal else "comment",
                body=payload.body.strip(), internal=internal)
    if not internal:
        if admin and t.created_by != user.id:
            if t.first_response_at is None:
                t.first_response_at = _now()
            if t.assignee_id is None:
                _event(db, t, user, "assign", old="", new=_uname(user))
                t.assignee_id, t.assignee_name = user.id, _uname(user)
            if t.status == "new":
                _set_status(db, t, "in_progress", None, body="Tự động: đã được phản hồi")
        elif t.created_by == user.id and t.status == "waiting_user":
            _set_status(db, t, "in_progress", None, body="Tự động: người gửi đã phản hồi")
    db.commit()
    db.refresh(ev)
    return _event_dict(ev)


@router.post("/{ticket_id}/attachments", status_code=201)
def upload_attachments(
    ticket_id: int,
    files: List[UploadFile] = File(...),
    event_id: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    t = _get_ticket(db, ticket_id, user)
    admin = _is_admin(user)
    if t.status == "closed":
        raise HTTPException(400, "Yêu cầu đã đóng")
    if len(files) > MAX_FILES_PER_REQUEST:
        raise HTTPException(400, f"Tối đa {MAX_FILES_PER_REQUEST} tệp mỗi lần tải lên")
    existing = db.query(func.count(FeedbackAttachment.id)).filter(
        FeedbackAttachment.ticket_id == t.id).scalar() or 0
    if existing + len(files) > MAX_FILES_PER_TICKET:
        raise HTTPException(400, f"Mỗi yêu cầu tối đa {MAX_FILES_PER_TICKET} tệp đính kèm")
    if event_id:
        ev = db.get(FeedbackEvent, event_id)
        if not ev or ev.ticket_id != t.id or (ev.user_id != user.id and not admin):
            raise HTTPException(400, "event_id không hợp lệ")

    saved, created = [], []
    try:
        for f in files:
            name = _safe_name(f.filename)
            ext = os.path.splitext(name)[1].lower()
            if ext not in ALLOWED_EXT:
                raise HTTPException(400, f"Loại tệp không được hỗ trợ: {name}")
            stored = uuid.uuid4().hex + ext
            path = os.path.join(UPLOAD_ROOT, stored)
            size = 0
            with open(path, "wb") as out:
                saved.append(path)
                while True:
                    chunk = f.file.read(1024 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    if size > MAX_FILE_BYTES:
                        raise HTTPException(413, f"{name} vượt quá {MAX_FILE_BYTES // 1024 // 1024} MB")
                    out.write(chunk)
            if size == 0:
                raise HTTPException(400, f"Tệp rỗng: {name}")
            a = FeedbackAttachment(
                ticket_id=t.id, event_id=event_id or None, filename=name, stored_name=stored,
                content_type=mimetypes.guess_type(name)[0] or "application/octet-stream",
                size=size, uploaded_by=user.id, uploaded_by_name=_uname(user))
            db.add(a)
            created.append(a)
        t.last_activity_at = _now()
        db.commit()
    except Exception:
        db.rollback()
        for p in saved:
            try:
                os.remove(p)
            except OSError:
                pass
        raise
    return [_att_dict(a) for a in created]


@router.post("/{ticket_id}/reopen")
def reopen_ticket(ticket_id: int, payload: Optional[ReasonPayload] = None,
                  db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    t = _get_ticket(db, ticket_id, user)
    if t.status not in DONE_STATUSES:
        raise HTTPException(400, "Chỉ mở lại được yêu cầu đã xử lý / đã đóng / bị từ chối")
    old, new = t.status, ("in_progress" if t.assignee_id else "new")
    t.status, t.resolved_at, t.closed_at = new, None, None
    t.reopen_count = (t.reopen_count or 0) + 1
    t.rating, t.rating_comment = None, None
    _event(db, t, user, "reopen", body=((payload.reason if payload else None) or "").strip() or None,
           old=old, new=new)
    db.commit()
    db.refresh(t)
    return _detail(t, user)


@router.post("/{ticket_id}/close")
def close_ticket(ticket_id: int, payload: Optional[ReasonPayload] = None,
                 db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    t = _get_ticket(db, ticket_id, user)
    if t.status == "closed":
        raise HTTPException(400, "Yêu cầu đã đóng")
    old = t.status
    t.status, t.closed_at = "closed", _now()
    t.resolved_at = t.resolved_at or t.closed_at
    _event(db, t, user, "close", body=((payload.reason if payload else None) or "").strip() or None,
           old=old, new="closed")
    db.commit()
    db.refresh(t)
    return _detail(t, user)


@router.post("/{ticket_id}/rate")
def rate_ticket(ticket_id: int, payload: RatePayload, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    t = _get_ticket(db, ticket_id, user)
    if t.created_by != user.id:
        raise HTTPException(403, "Chỉ người gửi mới được đánh giá")
    if t.status not in ("resolved", "closed"):
        raise HTTPException(400, "Chỉ đánh giá được khi yêu cầu đã xử lý / đã đóng")
    t.rating, t.rating_comment = payload.rating, (payload.comment or "").strip() or None
    _event(db, t, user, "rate", body=t.rating_comment, new=str(payload.rating))
    db.commit()
    db.refresh(t)
    return _detail(t, user)


@router.delete("/{ticket_id}")
def delete_ticket(ticket_id: int, db: Session = Depends(get_db),
                  user: User = Depends(require_admin)):
    t = db.get(_T, ticket_id)
    if not t:
        raise HTTPException(404, "Không tìm thấy yêu cầu")
    files = [a.stored_name for a in t.attachments]
    db.delete(t)
    db.commit()
    for fn in files:
        try:
            os.remove(os.path.join(UPLOAD_ROOT, fn))
        except OSError:
            pass
    try:
        log_action(db, user, "DELETE", "feedback_tickets", ticket_id)
    except Exception:
        pass
    return {"message": "Deleted"}
